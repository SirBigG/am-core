from io import BytesIO
from tempfile import NamedTemporaryFile

from django.contrib import admin
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import Workbook

from core.registry.admin import VarietyAdmin
from core.registry.forms import RegistryImportForm
from core.registry.models import Company, RegistryImportJob, Variety, VarietyCategory
from core.registry.parser import import_registry_workbook, parse_registry_date
from core.registry.parser_row_types import (
    ACTIVE_REGISTRY_ITEM_COLUMN_COUNT,
    INACTIVE_REGISTRY_ITEM_COLUMN_COUNT,
    normalize_active_registry_row,
    normalize_inactive_registry_row,
)
from core.utils.tests.factories import CountryFactory, UserFactory


def make_active_registry_row(**overrides):
    values = {
        "base_category_title": "Плодові та ягідні",
        "base_category_title_en": "Fruit and Berry",
        "child_category_title": "Абрикос звичайний",
        "child_category_title_en": "Apricot",
        "child_category_title_lt": "Prunus armeniaca L.",
        "application_number": "24256002",
        "date": "14.10.2024",
        "title": "Бержерон 1",
        "title_original": "Bergeron 1",
        "title_translit": "Bergeron 1",
        "selection_number": "Bergeron 1",
        "registration_year": 2024,
        "certificate_number": "240932",
        "patent": "®",
        "creation_method": "",
        "recommended_zone": "СЛП",
        "direction_of_use": "універс.",
        "ripeness_group": "сер.",
        "quality": "",
        "variety_description": "Опис сорту",
        "eu_upov_description": "",
        "variety_status": "",
        "original_country": "UA",
        "registration_country": "UA",
        "applicant": 3218,
        "applicant2": 3241,
        "applicant3": None,
        "applicant4": None,
        "applicant5": None,
        "applicant6": None,
        "owner": 3218,
        "owner2": 3241,
        "owner3": None,
        "owner4": None,
        "owner5": None,
        "owner6": None,
        "breeder": None,
        "breeder2": None,
        "breeder3": None,
        "breeder4": None,
        "breeder5": None,
        "breeder6": None,
    }
    values.update(overrides)
    field_order = [
        "base_category_title",
        "base_category_title_en",
        "child_category_title",
        "child_category_title_en",
        "child_category_title_lt",
        "application_number",
        "date",
        "title",
        "title_original",
        "title_translit",
        "selection_number",
        "registration_year",
        "certificate_number",
        "patent",
        "creation_method",
        "recommended_zone",
        "direction_of_use",
        "ripeness_group",
        "quality",
        "variety_description",
        "eu_upov_description",
        "variety_status",
        "original_country",
        "registration_country",
        "applicant",
        "applicant2",
        "applicant3",
        "applicant4",
        "applicant5",
        "applicant6",
        "owner",
        "owner2",
        "owner3",
        "owner4",
        "owner5",
        "owner6",
        "breeder",
        "breeder2",
        "breeder3",
        "breeder4",
        "breeder5",
        "breeder6",
    ]
    return [values[field] for field in field_order]


def make_registry_workbook_upload(include_invalid_inactive_row=False):
    workbook = Workbook()
    active_sheet = workbook.active
    active_sheet.title = "Реєстр_29_06_2026"
    active_sheet.append([f"Column {index}" for index in range(42)])
    active_sheet.append(make_active_registry_row())

    inactive_sheet = workbook.create_sheet("Виключені_з_реєстру")
    inactive_sheet.append([f"Column {index}" for index in range(43)])
    if include_invalid_inactive_row:
        inactive_sheet.append(["not-a-date"] + make_active_registry_row(title="Невалідний рядок"))

    for title in ["заявники", "власники", "підтримувачі"]:
        sheet = workbook.create_sheet(title)
        sheet.append(
            [
                "Номер заявки",
                "Порядок подачі",
                "Код",
                "Найменування українською",
                "Найменування оригінальне",
                "Код країни",
            ]
        )
        sheet.append(["24256002", 1, 3218, "Тестова компанія", "Test Company", "UA"])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return SimpleUploadedFile(
        "registry.xlsx",
        output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def save_registry_workbook(path, include_invalid_inactive_row=False):
    upload = make_registry_workbook_upload(include_invalid_inactive_row=include_invalid_inactive_row)
    with open(path, "wb") as workbook_file:
        workbook_file.write(upload.read())


class RegistryPublicViewTests(TestCase):
    def setUp(self):
        self.root = VarietyCategory.objects.create(title="Зернові", slug="grain")
        self.child = VarietyCategory.objects.create(title="Пшениця", slug="wheat", parent=self.root)

    def test_registry_index_renders_root_categories(self):
        response = self.client.get(reverse("registry:index"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "registry/index.html")
        self.assertIn(self.root, response.context["object_list"])

    def test_registry_category_renders_child_categories(self):
        response = self.client.get(reverse("registry:index-parent", kwargs={"root_slug": self.root.slug}))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "registry/categories.html")
        self.assertEqual(response.context["category"], self.root)
        self.assertIn(self.child, response.context["object_list"])

    def test_registry_variety_list_renders_child_varieties(self):
        country = CountryFactory(short_slug="ua")
        Variety.objects.create(title="Сорт А", slug="sort-a", category=self.child, original_country=country)

        response = self.client.get(
            reverse("registry:index-parent", kwargs={"root_slug": self.root.slug, "child_slug": self.child.slug})
        )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "registry/varieties.html")
        self.assertEqual(response.context["category"], self.child)
        self.assertEqual(response.context["posts"][0][0], "С")

    def test_registry_variety_list_does_not_write_during_get(self):
        Variety.objects.create(title="Сорт А", slug="sort-a", category=self.child)

        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(
                reverse("registry:index-parent", kwargs={"root_slug": self.root.slug, "child_slug": self.child.slug})
            )

        write_queries = [
            query["sql"]
            for query in queries
            if query["sql"].lstrip().upper().startswith(("UPDATE ", "INSERT ", "DELETE "))
        ]
        self.assertEqual(response.status_code, 200)
        self.assertEqual(write_queries, [])

    def test_registry_variety_list_returns_404_for_unknown_child(self):
        response = self.client.get(
            reverse("registry:index-parent", kwargs={"root_slug": self.root.slug, "child_slug": "missing"})
        )

        self.assertEqual(response.status_code, 404)


class RegistryParserRowTests(TestCase):
    def test_normalize_active_registry_row_accepts_new_2026_layout(self):
        row = make_active_registry_row()

        normalized = normalize_active_registry_row(row)

        self.assertEqual(len(normalized), ACTIVE_REGISTRY_ITEM_COLUMN_COUNT)
        self.assertEqual(normalized[10], "Bergeron 1")
        self.assertEqual(normalized[11], 2024)
        self.assertEqual(normalized[22], "UA")
        self.assertEqual(normalized[24], 3218)

    def test_normalize_active_registry_row_preserves_old_layout_alignment(self):
        new_row = make_active_registry_row()
        old_row = new_row[:10] + new_row[11:12] + new_row[13:19] + new_row[21:]

        normalized = normalize_active_registry_row(old_row)

        self.assertEqual(len(normalized), ACTIVE_REGISTRY_ITEM_COLUMN_COUNT)
        self.assertIsNone(normalized[10])
        self.assertEqual(normalized[11], 2024)
        self.assertIsNone(normalized[12])
        self.assertIsNone(normalized[19])
        self.assertIsNone(normalized[20])
        self.assertEqual(normalized[22], "UA")
        self.assertEqual(normalized[24], 3218)

    def test_normalize_inactive_registry_row_accepts_new_2026_layout(self):
        row = ["04.01.2016"] + make_active_registry_row()

        normalized = normalize_inactive_registry_row(row)

        self.assertEqual(len(normalized), INACTIVE_REGISTRY_ITEM_COLUMN_COUNT)
        self.assertEqual(normalized[0], "04.01.2016")
        self.assertEqual(normalized[12], 2024)
        self.assertEqual(normalized[25], 3218)

    def test_parse_registry_date_accepts_iso_ods_date_values(self):
        parsed_date, parsed_year = parse_registry_date("2016-01-04")

        self.assertEqual(parsed_date.isoformat(), "2016-01-04")
        self.assertEqual(parsed_year, 2016)


class RegistryImportTests(TestCase):
    def setUp(self):
        self.country = CountryFactory(short_slug="ua")
        self.applicant = Company.objects.create(code=3218, name="Старий заявник", country=self.country)
        self.owner = Company.objects.create(code=3241, name="Старий власник", country=self.country)

    def test_active_variety_import_updates_existing_without_deleting(self):
        category = VarietyCategory.objects.create(title="Абрикос звичайний")
        variety = Variety.objects.create(
            title="Бержерон 1",
            title_original="Old",
            slug="bergeron-1",
            category=category,
            excluded=True,
            unregister_year=2020,
        )

        imported = Variety.save_active_variety_from_row(make_active_registry_row(title_original="Bergeron 1 Updated"))

        variety.refresh_from_db()
        self.assertEqual(imported.pk, variety.pk)
        self.assertEqual(Variety.objects.filter(title="Бержерон 1").count(), 1)
        self.assertEqual(variety.title_original, "Bergeron 1 Updated")
        self.assertEqual(variety.registration_year, 2024)
        self.assertEqual(variety.applicant_id, self.applicant.id)
        self.assertFalse(variety.excluded)
        self.assertIsNone(variety.unregister_year)

    def test_inactive_variety_import_marks_existing_excluded(self):
        category = VarietyCategory.objects.create(title="Абрикос звичайний")
        variety = Variety.objects.create(title="Бержерон 1", slug="bergeron-1", category=category)
        row = ["04.01.2016"] + make_active_registry_row()

        imported = Variety.save_inactive_variety_from_row(row)

        variety.refresh_from_db()
        self.assertEqual(imported.pk, variety.pk)
        self.assertTrue(variety.excluded)
        self.assertEqual(variety.unregister_year, 2016)
        self.assertEqual(variety.unregister_date.isoformat(), "2016-01-04")


class RegistryAdminTests(TestCase):
    def setUp(self):
        self.user = UserFactory(is_staff=True, is_superuser=True)
        self.client.force_login(self.user)
        CountryFactory(short_slug="ua")

    def test_variety_admin_displays_original_country_as_fifth_column(self):
        model_admin = VarietyAdmin(Variety, admin.site)

        self.assertEqual(model_admin.list_display[4], "original_country")

    def test_registry_import_form_has_no_admin_run_option(self):
        form = RegistryImportForm()

        self.assertNotIn("run_now", form.fields)

    def test_variety_admin_import_page_queues_uploaded_workbook(self):
        response = self.client.post(
            reverse("admin:registry_variety_import_registry"),
            {"registry_file": make_registry_workbook_upload()},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        job = RegistryImportJob.objects.get()
        self.assertEqual(job.status, RegistryImportJob.Status.PENDING)
        self.assertEqual(job.original_filename, "registry.xlsx")
        self.assertEqual(job.created_by, self.user)
        self.assertFalse(Variety.objects.filter(title="Бержерон 1").exists())

    def test_registry_import_job_admin_has_no_run_actions(self):
        job = RegistryImportJob.objects.create(
            source_file=make_registry_workbook_upload(),
            original_filename="registry.xlsx",
            created_by=self.user,
        )

        changelist_response = self.client.get(reverse("admin:registry_registryimportjob_changelist"))
        change_response = self.client.get(reverse("admin:registry_registryimportjob_change", args=[job.id]))

        self.assertEqual(changelist_response.status_code, 200)
        self.assertEqual(change_response.status_code, 200)
        self.assertNotContains(changelist_response, "run_selected_jobs")
        self.assertNotContains(change_response, "Start import")

    def test_registry_import_job_command_processes_queued_upload(self):
        job = RegistryImportJob.objects.create(
            source_file=make_registry_workbook_upload(),
            original_filename="registry.xlsx",
            created_by=self.user,
        )

        call_command("run_registry_import_jobs")

        job.refresh_from_db()
        self.assertEqual(job.status, RegistryImportJob.Status.SUCCEEDED)
        self.assertEqual(job.summary["active_varieties"], 1)
        self.assertTrue(Company.objects.filter(code=3218, name="Тестова компанія").exists())
        variety = Variety.objects.get(title="Бержерон 1")
        self.assertEqual(variety.original_country.short_slug, "ua")
        self.assertEqual(variety.registration_country.short_slug, "ua")
        self.assertEqual(variety.application_number, "24256002")

    def test_registry_import_job_command_records_failure(self):
        job = RegistryImportJob.objects.create(
            source_file=make_registry_workbook_upload(include_invalid_inactive_row=True),
            original_filename="registry.xlsx",
            created_by=self.user,
        )

        call_command("run_registry_import_jobs")

        job.refresh_from_db()
        self.assertEqual(job.status, RegistryImportJob.Status.FAILED)
        self.assertIn("ValueError", job.error)
        self.assertFalse(Company.objects.filter(code=3218).exists())
        self.assertFalse(Variety.objects.filter(title="Бержерон 1").exists())

    def test_registry_import_job_command_does_not_claim_pending_job_while_another_is_running(self):
        RegistryImportJob.objects.create(
            source_file=make_registry_workbook_upload(),
            original_filename="running.xlsx",
            status=RegistryImportJob.Status.RUNNING,
            created_by=self.user,
        )
        pending_job = RegistryImportJob.objects.create(
            source_file=make_registry_workbook_upload(),
            original_filename="pending.xlsx",
            created_by=self.user,
        )

        call_command("run_registry_import_jobs")

        pending_job.refresh_from_db()
        self.assertEqual(pending_job.status, RegistryImportJob.Status.PENDING)
        self.assertFalse(Variety.objects.filter(title="Бержерон 1").exists())

    def test_registry_import_returns_company_counts(self):
        with NamedTemporaryFile(suffix=".xlsx") as workbook_file:
            save_registry_workbook(workbook_file.name)

            result = import_registry_workbook(workbook_file.name)

        self.assertEqual(result["applicants"], 1)
        self.assertEqual(result["owners"], 1)
        self.assertEqual(result["breeders"], 1)
        self.assertEqual(result["active_varieties"], 1)
        self.assertEqual(result["inactive_varieties"], 0)

    def test_registry_import_rolls_back_when_later_sheet_fails(self):
        with NamedTemporaryFile(suffix=".xlsx") as workbook_file:
            save_registry_workbook(workbook_file.name, include_invalid_inactive_row=True)

            with self.assertRaises(ValueError):
                import_registry_workbook(workbook_file.name)

        self.assertFalse(Company.objects.filter(code=3218).exists())
        self.assertFalse(Variety.objects.filter(title="Бержерон 1").exists())
